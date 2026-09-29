import io
from html import escape

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def generate_word_report(
	ai_summary_text: str,
	anomaly_df: pd.DataFrame,
	dataset_name: str = "Uploaded Dataset",
) -> bytes:
	"""Generate a crisp, executive-ready Word document report."""
	doc = Document()
	title = doc.add_heading(
		"Infesights AI — Data & Anomaly Intelligence Report",
		level=0,
	)
	title.alignment = WD_ALIGN_PARAGRAPH.CENTER

	metadata = doc.add_paragraph()
	metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
	metadata_run = metadata.add_run(
		f"Dataset Analyzed: {dataset_name}  |  "
		f"Total Anomalies Flagged: {len(anomaly_df)}"
	)
	metadata_run.font.italic = True
	metadata_run.font.color.rgb = RGBColor(100, 100, 100)

	doc.add_paragraph()
	doc.add_heading("📌 Executive AI Summary", level=1)
	if ai_summary_text and "No AI anomaly summary" not in ai_summary_text:
		doc.add_paragraph(ai_summary_text)
	else:
		doc.add_paragraph(
			f"The Infesights AI engine scanned {dataset_name} and identified "
			f"{len(anomaly_df)} statistical anomalies where key operational metrics "
			"deviated significantly from historical baselines. Review the flagged "
			"events below for detailed numbers and recommended action items."
		)

	doc.add_paragraph()
	doc.add_heading("🚨 Flagged Anomaly Events", level=1)
	if not anomaly_df.empty:
		table = doc.add_table(rows=1, cols=5)
		table.style = "Table Grid"
		headers = [
			"Timestamp",
			"Metric Name",
			"Observed Value",
			"Expected Baseline",
			"Confidence",
		]
		for index, header in enumerate(headers):
			cell = table.rows[0].cells[index]
			cell.text = header
			for paragraph in cell.paragraphs:
				for run in paragraph.runs:
					run.font.bold = True

		for _, row in anomaly_df.iterrows():
			cells = table.add_row().cells
			cells[0].text = str(row.get("timestamp", ""))[:19]
			cells[1].text = str(row.get("metric_name", row.get("feature", "N/A")))
			try:
				observed = f"{float(row.get('anomalous_value', 0)):,.2f}"
				baseline = f"{float(row.get('baseline_mean', 0)):,.2f}"
				confidence = f"{float(row.get('confidence_score', 0)):.1f}%"
			except (TypeError, ValueError):
				observed = str(row.get("anomalous_value", ""))
				baseline = str(row.get("baseline_mean", ""))
				confidence = str(row.get("confidence_score", ""))
			cells[2].text = observed
			cells[3].text = baseline
			cells[4].text = confidence
	else:
		doc.add_paragraph("✅ No critical statistical anomalies were detected in this dataset.")

	doc.add_heading("💡 Recommended Next Steps", level=1)
	steps = doc.add_paragraph()
	steps.add_run("1. Priority Audit: ").bold = True
	steps.add_run("Investigate top extreme spikes highlighted in the table above.\n")
	steps.add_run("2. Root Cause Check: ").bold = True
	steps.add_run(
		"Cross-reference metric timestamps against system deployment logs or "
		"operational events.\n"
	)
	steps.add_run("3. Data Hygiene: ").bold = True
	steps.add_run(
		"Verify that unformatted data inputs or zero-value records did not trigger "
		"false positive variance."
	)

	doc_io = io.BytesIO()
	doc.save(doc_io)
	return doc_io.getvalue()


def generate_excel_report(full_df: pd.DataFrame, anomaly_df: pd.DataFrame) -> bytes:
	"""Generate a clean dual-sheet Excel workbook."""
	output = io.BytesIO()
	with pd.ExcelWriter(output, engine="openpyxl") as writer:
		if not anomaly_df.empty:
			clean_anomalies = anomaly_df.copy()
			for column in clean_anomalies.select_dtypes(include=["float"]).columns:
				clean_anomalies[column] = clean_anomalies[column].round(2)
			clean_anomalies.to_excel(
				writer,
				sheet_name="Flagged Anomalies",
				index=False,
			)
		if full_df is not None and not full_df.empty:
			full_df.to_excel(writer, sheet_name="Full Raw Data", index=False)
	return output.getvalue()


def create_executive_pdf(report: str) -> bytes:
	buffer = io.BytesIO()
	document = SimpleDocTemplate(
		buffer,
		pagesize=letter,
		leftMargin=0.75 * inch,
		rightMargin=0.75 * inch,
		topMargin=0.7 * inch,
		bottomMargin=0.7 * inch,
		title="Executive Anomaly Post-Mortem",
	)
	styles = getSampleStyleSheet()
	styles.add(
		ParagraphStyle(
			name="ReportHeading",
			parent=styles["Heading2"],
			textColor=colors.HexColor("#17324D"),
			spaceBefore=12,
			spaceAfter=6,
		)
	)
	styles.add(
		ParagraphStyle(
			name="ReportBody",
			parent=styles["BodyText"],
			fontName="Helvetica",
			fontSize=9.5,
			leading=14,
			alignment=TA_LEFT,
			spaceAfter=6,
		)
	)
	story = [
		Paragraph("Executive Anomaly Post-Mortem", styles["Title"]),
		Spacer(1, 12),
	]
	for line in report.splitlines():
		plain_text = line.replace("**", "").replace("`", "")
		plain_text = plain_text.encode("ascii", "replace").decode("ascii")
		if not plain_text.strip():
			story.append(Spacer(1, 5))
		elif plain_text.startswith("### "):
			story.append(Paragraph(escape(plain_text[4:]), styles["ReportHeading"]))
		elif plain_text.startswith("#### "):
			story.append(Paragraph(escape(plain_text[5:]), styles["Heading3"]))
		else:
			story.append(Paragraph(escape(plain_text), styles["ReportBody"]))
	document.build(story)
	return buffer.getvalue()